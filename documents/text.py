"""Text extraction with OCR fallback.

PDF pages with a usable text layer are read directly. Pages without one (scans) are
sent to an OCR engine: macOS Vision (offline, via bin/vision_ocr) or Tesseract when
installed. If no engine is available the document is still stored and the status
says so; manual entry is never blocked.
"""
import json
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from django.conf import settings

MIN_TEXT_CHARS = 25
VISION_BIN = Path(settings.BASE_DIR) / "bin" / "vision_ocr"


@dataclass
class TextResult:
    pages: list = field(default_factory=list)
    ocr_pages: list = field(default_factory=list)
    status: str = "text_layer"
    engine: str = ""
    error: str = ""


def ocr_engine():
    choice = settings.OCR_ENGINE
    if choice in {"auto", "vision"} and VISION_BIN.exists():
        return "vision"
    if choice in {"auto", "tesseract"} and shutil.which("tesseract"):
        return "tesseract"
    return None


def ocr_engine_label():
    engine = ocr_engine()
    return {"vision": "macOS Vision (offline)", "tesseract": "Tesseract (offline)"}.get(engine, "Not available")


def _vision(path, pages=None):
    args = [str(VISION_BIN), str(path)]
    if pages:
        args.append(",".join(str(p) for p in pages))
    out = subprocess.run(args, capture_output=True, timeout=600, check=True)
    return json.loads(out.stdout.decode("utf-8"))


def _tesseract_image(path):
    out = subprocess.run(["tesseract", str(path), "-", "--psm", "3"], capture_output=True, timeout=300, check=True)
    return out.stdout.decode("utf-8", errors="replace")


def _tesseract_pdf_pages(path, pages):
    """Rasterize with pdftoppm if present (poppler), then OCR each page."""
    import tempfile

    if not shutil.which("pdftoppm"):
        raise RuntimeError("Tesseract needs pdftoppm (poppler) to OCR PDF pages.")
    texts = {}
    with tempfile.TemporaryDirectory() as tmp:
        for p in pages:
            prefix = Path(tmp) / f"p{p}"
            subprocess.run(["pdftoppm", "-r", "300", "-f", str(p), "-l", str(p), "-png", str(path), str(prefix)],
                           check=True, timeout=300)
            png = next(Path(tmp).glob(f"p{p}*.png"))
            texts[p] = _tesseract_image(png)
    return texts


def extract_text(path, content_type):
    path = Path(path)
    if content_type == "application/pdf":
        return _pdf(path)
    if content_type.startswith("image/"):
        return _image(path)
    if content_type.endswith("wordprocessingml.document"):
        return _docx(path)
    if content_type.startswith("text/"):
        return TextResult(pages=[path.read_text(errors="replace")], status="text_layer")
    return TextResult(status="not_applicable")


def _pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - damaged page; OCR may still work
            pages.append("")
    needs_ocr = [i + 1 for i, t in enumerate(pages) if len(t.strip()) < MIN_TEXT_CHARS]
    result = TextResult(pages=pages, status="text_layer")
    if not needs_ocr:
        return result
    engine = ocr_engine()
    if engine is None:
        result.status = "ocr_unavailable"
        result.error = f"Pages {needs_ocr} have no text layer and no OCR engine is available."
        return result
    try:
        if engine == "vision":
            ocr = _vision(path, needs_ocr)
            texts = {p: ocr[p - 1] for p in needs_ocr if p - 1 < len(ocr)}
        else:
            texts = _tesseract_pdf_pages(path, needs_ocr)
    except Exception as exc:  # noqa: BLE001
        result.status = "failed" if len(needs_ocr) == len(pages) else "partial_ocr"
        result.error = f"OCR failed: {exc}"
        return result
    for p, text in texts.items():
        pages[p - 1] = text
    result.ocr_pages = sorted(texts)
    result.engine = engine
    result.status = "ocr_complete" if len(needs_ocr) == len(pages) else "partial_ocr"
    return result


def _image(path):
    engine = ocr_engine()
    if engine is None:
        return TextResult(pages=[""], status="ocr_unavailable", error="No OCR engine is available.")
    try:
        text = _vision(path)[0] if engine == "vision" else _tesseract_image(path)
    except Exception as exc:  # noqa: BLE001
        return TextResult(pages=[""], status="failed", error=f"OCR failed: {exc}")
    return TextResult(pages=[text], ocr_pages=[1], status="ocr_complete", engine=engine)


def _docx(path):
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    # DOCX has no fixed pagination; the whole body is treated as page 1 and
    # references use section headings instead.
    return TextResult(pages=["\n".join(parts)], status="text_layer")


def page_count(path, content_type):
    if content_type == "application/pdf":
        from pypdf import PdfReader

        try:
            return len(PdfReader(str(path)).pages)
        except Exception:  # noqa: BLE001
            return None
    return 1
